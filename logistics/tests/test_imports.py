import io
import json
from decimal import Decimal
from unittest import mock

import pytest
from openpyxl import Workbook, load_workbook

from logistics import imports, services
from logistics.imports import apply_manifest, build_template, parse_manifest
from logistics.models import Cargo, Client
from logistics.services import InvalidTransition

pytestmark = pytest.mark.django_db

HEADERS = ["артикул", "телефон", "имя", "описание", "мест", "вес", "объём", "сумма"]
OPTIONAL_HEADERS = ["трек", "продавец", "упаковка", "хрупкое", "заметка"]


def xlsx(rows, headers=HEADERS):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "manifest.xlsx"
    return buf


@pytest.fixture
def truck():
    return services.create_truck()


@pytest.fixture
def ivan():
    client, _ = services.get_or_create_client("+79141234567", "Иван")
    return client


def test_valid_file_parses_rows(ivan):
    preview = parse_manifest(
        xlsx(
            [
                ["A-1", "8 914 123 45 67", "Иван", "Куртки", 2, 12.5, 0.3, 1500],
                ["A-2", "+7 (914) 765-43-21", "Пётр", "", 1, None, None, None],
            ]
        )
    )

    assert preview.errors == []
    assert not preview.has_errors
    assert len(preview.rows) == 2
    first, second = preview.rows
    assert first.line == 2
    assert first.article == "A-1"
    assert first.phone == "+79141234567"
    assert first.places == 2
    assert first.weight_kg == Decimal("12.50")
    assert first.volume_m3 == Decimal("0.300")
    assert first.amount_due == Decimal("1500.00")
    assert first.new_client is False
    assert second.phone == "+79147654321"
    assert second.places == 1
    assert second.weight_kg is None
    assert second.amount_due == Decimal("0")
    assert second.new_client is True
    assert preview.new_client_count == 1
    assert preview.error_count == 0


def test_header_variants_and_order():
    headers = ["  Телефон ", "АРТИКУЛ", "Объем", "Сумма ", "Мест"]
    preview = parse_manifest(xlsx([["89141234567", "Z9", "1,5", "100", 3]], headers=headers))

    assert preview.errors == []
    row = preview.rows[0]
    assert row.article == "Z9"
    assert row.volume_m3 == Decimal("1.500")
    assert row.amount_due == Decimal("100.00")
    assert row.places == 3
    assert row.description == ""
    assert row.name == ""


def test_comma_decimals_and_spaces():
    preview = parse_manifest(xlsx([["A", "89141234567", "", "", "2", "12,5", "0,25", "1 200,50"]]))

    row = preview.rows[0]
    assert row.errors == []
    assert row.weight_kg == Decimal("12.50")
    assert row.volume_m3 == Decimal("0.250")
    assert row.amount_due == Decimal("1200.50")


def test_numeric_phone_cell_is_accepted():
    preview = parse_manifest(xlsx([["A", 89141234567, "", "", 1, None, None, None]]))

    assert preview.rows[0].phone == "+79141234567"
    assert preview.rows[0].errors == []


@pytest.mark.parametrize(
    ("row", "message"),
    [
        (["", "89141234567", "", "", 1, None, None, None], "Нет артикула"),
        (["A", "", "", "", 1, None, None, None], "Нет телефона"),
        (["A", "12345", "", "", 1, None, None, None], "Неверный телефон"),
        (["A", "89141234567", "", "", 0, None, None, None], "Мест должно быть целым числом от 1"),
        (["A", "89141234567", "", "", 2.5, None, None, None], "Мест должно быть целым числом от 1"),
        (["A", "89141234567", "", "", "две", None, None, None], "Мест должно быть целым числом"),
        (["A", "89141234567", "", "", 1, "тяжёлый", None, None], "Вес: не число"),
        (["A", "89141234567", "", "", 1, None, "1,2,3", None], "Объём: не число"),
        (["A", "89141234567", "", "", 1, None, None, "-5"], "Сумма: не может быть меньше 0"),
        (["A" * 65, "89141234567", "", "", 1, None, None, None], "Артикул длиннее 64 символов"),
    ],
)
def test_row_errors(row, message):
    preview = parse_manifest(xlsx([row]))

    assert preview.has_errors
    assert preview.error_count == 1
    assert any(message in error for error in preview.rows[0].errors), preview.rows[0].errors


def test_duplicate_article_within_file():
    preview = parse_manifest(
        xlsx(
            [
                ["A-1", "89141234567", "", "", 1, None, None, None],
                ["a-1 ", "89141234568", "", "", 1, None, None, None],
            ]
        )
    )

    assert preview.rows[0].errors == []
    assert preview.rows[1].errors == ["Артикул повторяется (строка 2)"]
    assert preview.error_count == 1


def test_empty_rows_are_skipped():
    preview = parse_manifest(
        xlsx([["A", "89141234567", "", "", 1, None, None, None], [None] * 8, ["", "", "", ""]])
    )

    assert len(preview.rows) == 1


def test_missing_required_header_is_file_error():
    preview = parse_manifest(xlsx([["A", "Иван"]], headers=["артикул", "имя"]))

    assert preview.has_errors
    assert preview.errors == ["Нет колонки «телефон»"]
    assert preview.rows == []


def test_not_an_xlsx_is_file_error():
    buf = io.BytesIO(b"definitely not a zip")
    buf.name = "manifest.xlsx"

    preview = parse_manifest(buf)

    assert preview.has_errors
    assert preview.errors == ["Не удалось прочитать файл. Сохраните его как .xlsx"]


def test_empty_sheet_is_file_error():
    preview = parse_manifest(xlsx([], headers=[]))

    assert preview.errors == ["В файле нет строк с грузами"]


def test_new_clients_counted_once_per_phone(ivan):
    preview = parse_manifest(
        xlsx(
            [
                ["A", "89140000001", "", "", 1, None, None, None],
                ["B", "8 914 000 00 01", "", "", 1, None, None, None],
                ["C", "89141234567", "", "", 1, None, None, None],
            ]
        )
    )

    assert preview.new_client_count == 1


def test_preview_roundtrips_through_json():
    preview = parse_manifest(xlsx([["A", "89141234567", "Иван", "x", 2, "1,5", None, 10]]))

    restored = imports.ImportPreview.from_dict(json.loads(json.dumps(preview.to_dict())))

    assert restored == preview


def test_apply_creates_clients_and_loads_cargo(truck, ivan):
    preview = parse_manifest(
        xlsx(
            [
                ["A-1", "89141234567", "Другое имя", "Куртки", 2, "12,5", None, 1500],
                ["A-2", "89147654321", "Пётр", "", 1, None, None, None],
            ]
        )
    )

    count = apply_manifest(preview, truck)

    assert count == 2
    ivan.refresh_from_db()
    assert ivan.name == "Иван"
    petr = Client.objects.get(phone="+79147654321")
    assert petr.name == "Пётр"
    first = Cargo.objects.get(article="A-1")
    assert first.client == ivan
    assert first.truck == truck
    assert first.status == Cargo.Status.LOADED
    assert first.places == 2
    assert first.weight_kg == Decimal("12.50")
    assert first.amount_due == Decimal("1500.00")


def test_apply_is_atomic(truck):
    preview = parse_manifest(
        xlsx(
            [
                ["A-1", "89140000001", "", "", 1, None, None, None],
                ["A-2", "89140000002", "", "", 1, None, None, None],
            ]
        )
    )
    real_add = services.add_cargo
    calls = []

    def flaky(*args, **kwargs):
        calls.append(1)
        if len(calls) == 2:
            raise RuntimeError("boom")
        return real_add(*args, **kwargs)

    with mock.patch.object(imports.services, "add_cargo", side_effect=flaky):
        with pytest.raises(RuntimeError):
            apply_manifest(preview, truck)

    assert Cargo.objects.count() == 0
    assert Client.objects.count() == 0


def test_apply_rejects_non_loading_truck(truck):
    services.truck_depart(truck)
    preview = parse_manifest(xlsx([["A-1", "89140000001", "", "", 1, None, None, None]]))

    with pytest.raises(InvalidTransition):
        apply_manifest(preview, truck)

    assert Cargo.objects.count() == 0
    assert Client.objects.count() == 0


def test_apply_refuses_preview_with_errors(truck):
    preview = parse_manifest(xlsx([["", "89140000001", "", "", 1, None, None, None]]))

    with pytest.raises(ValueError):
        apply_manifest(preview, truck)

    assert Cargo.objects.count() == 0


def test_template_has_headers():
    wb = load_workbook(io.BytesIO(build_template()))

    assert [cell.value for cell in wb.active[1]] == [*HEADERS, *OPTIONAL_HEADERS]


# ---------- optional columns: трек, продавец, упаковка, хрупкое, заметка ----------


def test_optional_columns_are_parsed():
    headers = [*HEADERS, "Трек", "ПРОДАВЕЦ", "Упаковка", "хрупкое", "Заметка"]
    rows = [
        ["A-1", "89141234567", "", "", 1, None, None, None, "YT1", "1688", "Мешок", "да", "Стекло"],
        ["A-2", "89141234567", "", "", 1, None, None, None, None, None, "pallet", True, None],
        ["A-3", "89141234567", "", "", 1, None, None, None, None, None, None, None, None],
    ]

    preview = parse_manifest(xlsx(rows, headers=headers))

    assert not preview.has_errors, [r.errors for r in preview.rows]
    first, second, third = preview.rows
    assert (first.track_number, first.seller, first.packaging, first.fragile, first.note) == (
        "YT1",
        "1688",
        "bag",
        True,
        "Стекло",
    )
    assert (second.packaging, second.fragile) == ("pallet", True)
    assert (third.track_number, third.packaging, third.fragile, third.note) == (
        "",
        "box",
        False,
        "",
    )


def test_file_without_optional_columns_uses_defaults():
    preview = parse_manifest(xlsx([["A-1", "89141234567", "", "", 1, None, None, None]]))

    row = preview.rows[0]
    assert (row.track_number, row.seller, row.packaging, row.fragile, row.note) == (
        "",
        "",
        "box",
        False,
        "",
    )


@pytest.mark.parametrize(
    "cells, message, attr",
    [
        ({"packaging": "ящик"}, "Упаковка", "packaging"),
        ({"fragile": "может быть"}, "Хрупкое", "fragile"),
        ({"track_number": "Y" * 65}, "Трек-номер длиннее 64 символов", "track_number"),
        ({"seller": "S" * 121}, "Продавец длиннее 120 символов", "seller"),
    ],
)
def test_optional_column_errors(cells, message, attr):
    preview = imports.parse_rows([{"article": "A", "phone": "89141234567", **cells}])

    row = preview.rows[0]
    assert any(message in error for error in row.errors), row.errors
    assert attr in row.invalid


def test_apply_saves_optional_fields(truck):
    preview = imports.parse_rows(
        [
            {
                "article": "A-1",
                "phone": "89141234567",
                "track_number": "YT1",
                "seller": "1688",
                "packaging": "other",
                "fragile": "да",
                "note": "Заметка",
            }
        ]
    )

    apply_manifest(preview, truck)

    cargo = Cargo.objects.get(article="A-1")
    assert (cargo.track_number, cargo.seller, cargo.packaging, cargo.fragile, cargo.note) == (
        "YT1",
        "1688",
        "other",
        True,
        "Заметка",
    )


# ---------- editable grid rows ----------


def test_parse_rows_numbers_rows_by_position_and_skips_empty_ones(ivan):
    rows = [
        {"article": "A-1", "phone": "8 914 123 45 67", "places": "2", "weight": "1,5"},
        {"article": "", "phone": "", "packaging": "box", "fragile": ""},  # an empty grid row
        {"article": "a-1", "phone": "89140000001", "amount": "abc"},
    ]

    preview = imports.parse_rows(rows)

    assert [row.line for row in preview.rows] == [1, 3]
    first, third = preview.rows
    assert first.errors == []
    assert first.places == 2 and first.weight_kg == Decimal("1.50")
    assert first.new_client is False
    assert "Артикул повторяется (строка 1)" in third.errors
    assert "Сумма: не число" in third.errors
    assert set(third.invalid) == {"article", "amount"}
    assert third.new_client is True
    assert preview.error_count == 1


def test_parse_rows_keeps_what_was_typed():
    preview = imports.parse_rows([{"article": " A-1 ", "phone": "12345", "places": "x"}])

    row = preview.rows[0]
    assert row.raw["phone"] == "12345"
    assert row.raw["places"] == "x"
    assert set(row.invalid) == {"phone", "places"}


def test_parse_rows_all_empty_is_an_error():
    preview = imports.parse_rows([{"article": " ", "packaging": "box"}, {}])

    assert preview.rows == []
    assert preview.errors == [imports.ROWS_EMPTY]
    assert preview.has_errors


def test_parse_rows_limits_row_count():
    rows = [{"article": f"A{n}", "phone": "89141234567"} for n in range(imports.MAX_ROWS + 1)]

    preview = imports.parse_rows(rows)

    assert preview.has_errors
    assert preview.errors and "1000" in preview.errors[0]


def test_grid_rows_from_file_preview_match_the_file():
    preview = parse_manifest(
        xlsx(
            [
                ["A-1", 89141234567, "Иван", "", 2.0, 12.5, None, 1500],
                ["A-2", "89140000001", "", "", 2.5, None, None, None],
            ]
        )
    )

    regrid = imports.parse_rows([row.raw for row in preview.rows])

    assert [r.errors for r in regrid.rows] == [r.errors for r in preview.rows]
    assert regrid.rows[0].weight_kg == Decimal("12.50")
    assert regrid.rows[0].places == 2
