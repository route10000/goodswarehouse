"""Truck manifest import from ``.xlsx`` and the editable entry grid (SPEC §10, §14.5).

``parse_manifest`` reads a workbook into an ``ImportPreview`` without writing anything;
``parse_rows`` validates the same columns typed into the panel's grid (one dict of cell
texts per row); ``apply_manifest`` loads a clean preview into a ``loading`` truck in one
transaction. The preview is plain data (``to_dict`` / ``from_dict``) so the panel can
keep it in the session between the preview and the confirmation instead of the file.
"""

import io
from dataclasses import asdict, dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from zipfile import BadZipFile

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException

from logistics import services
from logistics.models import Cargo, Client, Truck
from logistics.phones import InvalidPhone, normalize_phone

TEMPLATE_HEADERS = [
    "артикул",
    "телефон",
    "имя",
    "описание",
    "мест",
    "вес",
    "объём",
    "сумма",
    # optional
    "трек",
    "продавец",
    "упаковка",
    "хрупкое",
    "заметка",
]
REQUIRED = ("артикул", "телефон")
# Normalized header → cell key. «объём» and «объем» both normalize to «объем».
COLUMNS = {
    "артикул": "article",
    "телефон": "phone",
    "имя": "name",
    "описание": "description",
    "мест": "places",
    "вес": "weight",
    "объем": "volume",
    "сумма": "amount",
    "трек": "track_number",
    "трек-номер": "track_number",
    "трек номер": "track_number",
    "продавец": "seller",
    "упаковка": "packaging",
    "хрупкое": "fragile",
    "заметка": "note",
}
# Cell keys in grid order (the grid posts one list per key).
CELL_KEYS = [
    "article",
    "phone",
    "name",
    "description",
    "places",
    "weight",
    "volume",
    "amount",
    "track_number",
    "seller",
    "packaging",
    "fragile",
    "note",
]
# Select cells always carry a value, so they do not make a grid row «filled».
SELECT_KEYS = ("packaging", "fragile")
PACKAGING = {
    "коробка": Cargo.Packaging.BOX,
    "мешок": Cargo.Packaging.BAG,
    "паллета": Cargo.Packaging.PALLET,
    "палета": Cargo.Packaging.PALLET,
    "другое": Cargo.Packaging.OTHER,
    **{value: value for value in Cargo.Packaging.values},
}
FRAGILE_YES = {"да", "д", "+", "1", "true", "yes", "y", "x", "х", "хрупкое", "хрупкий"}
FRAGILE_NO = {"", "нет", "н", "-", "0", "false", "no", "n"}
MAX_ROWS = 1000
DECIMALS = {
    # attribute: (label, model field)
    "weight": ("Вес", "weight_kg"),
    "volume": ("Объём", "volume_m3"),
    "amount": ("Сумма", "amount_due"),
}
NBSP = " "
FILE_UNREADABLE = "Не удалось прочитать файл. Сохраните его как .xlsx"
FILE_EMPTY = "В файле нет строк с грузами"
ROWS_EMPTY = "Заполните хотя бы одну строку: артикул и телефон"
TOO_MANY_ROWS = f"Не больше {MAX_ROWS} строк за раз. Разбейте манифест на части"


@dataclass
class ManifestRow:
    line: int
    article: str = ""
    phone: str = ""
    name: str = ""
    description: str = ""
    places: int | None = 1
    weight_kg: Decimal | None = None
    volume_m3: Decimal | None = None
    amount_due: Decimal | None = Decimal("0")
    track_number: str = ""
    seller: str = ""
    packaging: str = Cargo.Packaging.BOX
    fragile: bool = False
    note: str = ""
    new_client: bool = False
    errors: list[str] = field(default_factory=list)
    invalid: list[str] = field(default_factory=list)  # cell keys with an error
    raw: dict = field(default_factory=dict)  # cell texts as typed, for the grid

    def fail(self, key: str, message: str) -> None:
        self.errors.append(message)
        if key not in self.invalid:
            self.invalid.append(key)


@dataclass
class ImportPreview:
    rows: list[ManifestRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)  # file-level problems

    @property
    def error_count(self) -> int:
        return sum(1 for row in self.rows if row.errors)

    @property
    def has_errors(self) -> bool:
        return bool(self.errors) or self.error_count > 0

    @property
    def new_client_count(self) -> int:
        return len({row.phone for row in self.rows if row.new_client})

    def to_dict(self) -> dict:
        def plain(value):
            return str(value) if isinstance(value, Decimal) else value

        return {
            "errors": list(self.errors),
            "rows": [{k: plain(v) for k, v in asdict(row).items()} for row in self.rows],
        }

    @classmethod
    def from_dict(cls, data: dict) -> "ImportPreview":
        rows = []
        for raw in data.get("rows", []):
            values = dict(raw)
            for name in ("weight_kg", "volume_m3", "amount_due"):
                if values.get(name) is not None:
                    values[name] = Decimal(values[name])
            rows.append(ManifestRow(**values))
        return cls(rows=rows, errors=list(data.get("errors", [])))


# ---------- cell parsing ----------


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace(NBSP, " ").strip()


def _header(value) -> str:
    return _text(value).lower().replace("ё", "е")


def _places(value) -> int:
    message = "Мест должно быть целым числом от 1"
    if value is None or _text(value) == "":
        return 1
    if isinstance(value, bool):
        raise ValueError(message)
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError(message)
        value = int(value)
    if isinstance(value, int):
        number = value
    else:
        try:
            number = int(_text(value).replace(" ", ""))
        except ValueError as exc:
            raise ValueError(message) from exc
    if number < 1:
        raise ValueError(message)
    return number


def _decimal(value, label: str, model_field: str) -> Decimal | None:
    if value is None or _text(value) == "":
        return None
    if isinstance(value, bool):
        raise ValueError(f"{label}: не число")
    text = _text(value).replace(" ", "").replace(",", ".")
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"{label}: не число") from exc
    if not number.is_finite():
        raise ValueError(f"{label}: не число")
    if number < 0:
        raise ValueError(f"{label}: не может быть меньше 0")
    db_field = Cargo._meta.get_field(model_field)
    number = number.quantize(Decimal(1).scaleb(-db_field.decimal_places), ROUND_HALF_UP)
    if len(number.as_tuple().digits) > db_field.max_digits:
        raise ValueError(f"{label}: слишком большое число")
    return number


def _packaging(text: str) -> str:
    key = text.strip().lower().replace("ё", "е")
    if not key:
        return Cargo.Packaging.BOX
    if key not in PACKAGING:
        raise ValueError("Упаковка: коробка, мешок, паллета или другое")
    return PACKAGING[key]


def _fragile(text: str) -> bool:
    key = text.strip().lower()
    if key in FRAGILE_YES:
        return True
    if key in FRAGILE_NO:
        return False
    raise ValueError("Хрупкое: да или нет")


def _max_length(model, name: str) -> int:
    return model._meta.get_field(name).max_length


def _is_empty(cells: dict) -> bool:
    return not any(_text(v) for k, v in cells.items() if k not in SELECT_KEYS)


def _parse_row(line: int, cells: dict) -> ManifestRow:
    raw = {key: _text(cells.get(key)) for key in CELL_KEYS}
    row = ManifestRow(line=line, raw=raw)
    row.article = raw["article"]
    row.name = raw["name"]
    row.description = raw["description"]
    row.track_number = raw["track_number"]
    row.seller = raw["seller"]
    row.note = raw["note"]
    raw_phone = raw["phone"]

    if not row.article:
        row.fail("article", "Нет артикула")
    elif len(row.article) > _max_length(Cargo, "article"):
        row.fail("article", "Артикул длиннее 64 символов")
    if not raw_phone:
        row.fail("phone", "Нет телефона")
    else:
        try:
            row.phone = normalize_phone(raw_phone)
        except InvalidPhone:
            row.phone = raw_phone
            row.fail("phone", f"Неверный телефон: {raw_phone}")
    if len(row.name) > _max_length(Client, "name"):
        row.fail("name", "Имя длиннее 200 символов")
    if len(row.description) > _max_length(Cargo, "description"):
        row.fail("description", "Описание длиннее 255 символов")
    if len(row.track_number) > _max_length(Cargo, "track_number"):
        row.fail("track_number", "Трек-номер длиннее 64 символов")
    if len(row.seller) > _max_length(Cargo, "seller"):
        row.fail("seller", "Продавец длиннее 120 символов")

    try:
        row.places = _places(raw["places"])
    except ValueError as exc:
        row.places = None
        row.fail("places", str(exc))
    for attr, (label, model_field) in DECIMALS.items():
        try:
            value = _decimal(raw[attr], label, model_field)
        except ValueError as exc:
            row.fail(attr, str(exc))
            value = None
        else:
            if value is None and model_field == "amount_due":
                value = Decimal("0")  # an empty amount means nothing to pay
        setattr(row, model_field, value)

    # Select cells: the grid shows the parsed choice; a bad value stays as typed.
    try:
        row.packaging = _packaging(raw["packaging"])
        raw["packaging"] = row.packaging
    except ValueError as exc:
        row.fail("packaging", str(exc))
    try:
        row.fragile = _fragile(raw["fragile"])
        raw["fragile"] = "да" if row.fragile else ""
    except ValueError as exc:
        row.fail("fragile", str(exc))
    return row


def _validate(entries) -> ImportPreview:
    """Parse ``(line, cells)`` pairs: per-row checks, repeated articles, new clients."""
    preview = ImportPreview()
    seen: dict[str, int] = {}
    for line, cells in entries:
        row = _parse_row(line, cells)
        if row.article:
            key = row.article.casefold()
            if key in seen:
                row.fail("article", f"Артикул повторяется (строка {seen[key]})")
            else:
                seen[key] = row.line
        preview.rows.append(row)

    valid_phones = {row.phone for row in preview.rows if row.phone.startswith("+")}
    known = set(Client.objects.filter(phone__in=valid_phones).values_list("phone", flat=True))
    for row in preview.rows:
        row.new_client = row.phone in valid_phones and row.phone not in known
    return preview


# ---------- public API ----------


def parse_manifest(file) -> ImportPreview:
    """Read an uploaded workbook. Never raises for bad content: errors go into the preview."""
    try:
        workbook = load_workbook(file, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, ValueError, OSError):
        return ImportPreview(errors=[FILE_UNREADABLE])
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()

    header_row = next((r for r in rows if any(_text(c) for c in r)), None)
    if header_row is None:
        return ImportPreview(errors=[FILE_EMPTY])
    header_index = rows.index(header_row)
    positions: dict[str, int] = {}
    for index, cell in enumerate(header_row):
        attr = COLUMNS.get(_header(cell))
        if attr and attr not in positions:
            positions[attr] = index
    missing = [name for name in REQUIRED if COLUMNS[name] not in positions]
    if missing:
        return ImportPreview(errors=[f"Нет колонки «{name}»" for name in missing])

    entries = []
    for offset, values in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if not any(_text(c) for c in values):
            continue
        cells = {
            attr: (values[index] if index < len(values) else None)
            for attr, index in positions.items()
        }
        entries.append((offset, cells))
    if not entries:
        return ImportPreview(errors=[FILE_EMPTY])
    if len(entries) > MAX_ROWS:
        return ImportPreview(errors=[TOO_MANY_ROWS])
    return _validate(entries)


def parse_rows(rows: list[dict]) -> ImportPreview:
    """Validate rows typed into the grid. ``line`` is the row's position (from 1);
    rows with nothing typed are skipped. Never raises for bad content."""
    entries = [(n, cells) for n, cells in enumerate(rows, start=1) if not _is_empty(cells)]
    if not entries:
        return ImportPreview(errors=[ROWS_EMPTY])
    if len(entries) > MAX_ROWS:
        return ImportPreview(errors=[TOO_MANY_ROWS])
    return _validate(entries)


def apply_manifest(preview: ImportPreview, truck: Truck) -> int:
    """Create clients and cargo for every row, all or nothing. Returns the cargo count."""
    if preview.has_errors:
        raise ValueError("Манифест с ошибками нельзя импортировать")
    with transaction.atomic():
        status = Truck.objects.select_for_update().values_list("status", flat=True).get(pk=truck.pk)
        if status != Truck.Status.LOADING:
            raise services.InvalidTransition(f"Машина {truck.number} уже не грузится")
        for row in preview.rows:
            client, _ = services.get_or_create_client(row.phone, row.name)
            services.add_cargo(
                client,
                row.article,
                description=row.description,
                places=row.places,
                weight_kg=row.weight_kg,
                volume_m3=row.volume_m3,
                amount_due=row.amount_due if row.amount_due is not None else Decimal("0"),
                track_number=row.track_number,
                seller=row.seller,
                packaging=row.packaging,
                fragile=row.fragile,
                note=row.note,
                truck=truck,
            )
    return len(preview.rows)


def build_template() -> bytes:
    """Empty manifest with the expected header row."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Манифест"
    sheet.append(TEMPLATE_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    widths = (18, 18, 22, 32, 8, 10, 10, 12, 20, 20, 12, 10, 32)
    for letter, width in zip("ABCDEFGHIJKLM", widths, strict=True):
        sheet.column_dimensions[letter].width = width
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
