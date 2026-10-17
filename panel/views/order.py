"""Ордер на день (SPEC §7): the call list screen, marks via HTMX, print view and Excel.

Lines and ContactAttempt are written only through ``logistics.orders``; this module
just reads, renders and turns ``OrderError`` into toasts. Both staff groups may build
the order and mark lines; adding and removing lines is for managers.
"""

from datetime import date as Date
from datetime import timedelta
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from logistics import orders
from logistics.models import Cargo, Client, DailyOrder, DailyOrderLine
from logistics.orders import OrderError
from panel.permissions import is_manager, manager_required

from .common import any_of, count_phrase, phone_q

CallStatus = DailyOrderLine.CallStatus
PICKED_UP = Cargo.Status.PICKED_UP
FILTERS = [
    ("", "Всего", "total"),
    (CallStatus.REACHED, "Позвонили", "reached"),
    (CallStatus.NO_ANSWER, "Не дозвонились", "no_answer"),
    (CallStatus.PENDING, "Осталось", "pending"),
]
CANDIDATE_LIMIT = 20


# ---------- helpers ----------


def _parse_day(day: str) -> Date:
    try:
        if len(day) != 10:
            raise ValueError
        return Date.fromisoformat(day)
    except ValueError:
        raise Http404 from None


def _day_url(day: Date) -> str:
    return reverse("panel:order_day", args=[day.isoformat()])


def remaining_lines(day: Date | None = None) -> int:
    """Lines nobody has called yet in the order for ``day`` (today); one query."""
    return (
        DailyOrderLine.objects.filter(
            order__date=day or timezone.localdate(), call_status=CallStatus.PENDING
        )
        .exclude(cargo__status=PICKED_UP)
        .count()
    )


def order_counts(order: DailyOrder | None) -> dict:
    if order is None:
        return {"total": 0, "reached": 0, "no_answer": 0, "pending": 0}
    return order.lines.aggregate(
        total=Count("id"),
        reached=Count("id", filter=Q(call_status=CallStatus.REACHED)),
        no_answer=Count("id", filter=Q(call_status=CallStatus.NO_ANSWER)),
        pending=Count("id", filter=Q(call_status=CallStatus.PENDING) & ~Q(cargo__status=PICKED_UP)),
    )


def _filters(order, status: str) -> list[dict]:
    counts = order_counts(order)
    return [
        {"value": value, "label": label, "count": counts[key], "current": value == status}
        for value, label, key in FILTERS
    ]


def numbered_lines(order: DailyOrder) -> list[DailyOrderLine]:
    """Lines in call order with a stable ``number``; picked-up cargo goes to the bottom."""
    lines = list(order.lines.select_related("cargo__client").order_by("id"))
    lines = [x for x in lines if not x.is_picked_up] + [x for x in lines if x.is_picked_up]
    for number, line in enumerate(lines, start=1):
        line.number = number
    return lines


def _matches(line: DailyOrderLine, status: str) -> bool:
    if not status:
        return True
    if status == CallStatus.PENDING and line.is_picked_up:
        return False
    return line.call_status == status


def _status_param(request) -> str:
    # The page has ?status=; HTMX requests from its rows send the same value as "filter".
    status = (
        request.GET.get("status") or request.GET.get("filter") or request.POST.get("filter") or ""
    )
    return status if status in CallStatus.values else ""


def _get_line(day: Date, pk: int) -> DailyOrderLine:
    line = get_object_or_404(
        DailyOrderLine.objects.select_related("order", "cargo__client"), pk=pk, order__date=day
    )
    for numbered in numbered_lines(line.order):
        if numbered.pk == line.pk:
            line.number = numbered.number
    return line


def _is_htmx(request) -> bool:
    return bool(request.headers.get("HX-Request"))


def _row_context(request, line, *, editing=False, toast="", error=False, deleted=False):
    day = line.order.date
    return {
        "line": line,
        "read_only": line.order.is_read_only,
        "manager": is_manager(request.user),
        "editing": editing and not line.order.is_read_only,
        "day": day,
        "order_url": _day_url(day),
        "filters": _filters(line.order, _status_param(request)),
        "remaining": remaining_lines(),
        "toast": toast,
        "error": error,
        "deleted": deleted,
        "oob": True,
    }


# ---------- page ----------


@login_required
def order_today(request):
    picked = request.GET.get("date")
    if picked:
        try:
            return redirect(_day_url(Date.fromisoformat(picked)))
        except ValueError:
            messages.error(request, "Неверная дата")
    return _render_order(request, timezone.localdate())


@login_required
def order_day(request, day):
    return _render_order(request, _parse_day(day))


def _render_order(request, day: Date):
    today = timezone.localdate()
    order = DailyOrder.objects.filter(date=day).first()
    status = _status_param(request)
    lines = [x for x in numbered_lines(order) if _matches(x, status)] if order else []
    return render(
        request,
        "panel/order/day.html",
        {
            "nav_active": "order",
            "day": day,
            "today": today,
            "is_today": day == today,
            "is_past": day < today,
            "prev_day": day - timedelta(days=1),
            "next_day": day + timedelta(days=1),
            "order": order,
            "read_only": day < today,
            "manager": is_manager(request.user),
            "status": status,
            "filters": _filters(order, status),
            "lines": lines,
            "order_url": _day_url(day),
        },
    )


# ---------- build / add (full page, toast after redirect) ----------


@require_POST
@login_required
def order_build(request, day):
    day = _parse_day(day)
    try:
        _order, added = orders.build_daily_order(day, request.user)
    except OrderError as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, f"Добавлено {count_phrase(added, 'строка', 'строки', 'строк')}")
    return redirect(_day_url(day))


@require_POST
@manager_required
def order_add(request, day):
    day = _parse_day(day)
    order = get_object_or_404(DailyOrder, date=day)
    cargo = get_object_or_404(Cargo, pk=request.POST.get("cargo") or 0)
    try:
        orders.add_line(order, cargo, request.user)
    except OrderError as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, f"Груз {cargo.article} добавлен в ордер")
    return redirect(_day_url(day))


@manager_required
def order_candidates(request, day):
    """Ready cargo not in the order yet, by article, phone or client code (HTMX)."""
    day = _parse_day(day)
    order = DailyOrder.objects.filter(date=day).first()
    q = request.GET.get("q", "").strip()
    cargos = []
    if len(q) >= 2:
        cargos = (
            Cargo.objects.filter(status=Cargo.Status.READY)
            .filter(
                any_of(
                    Q(article__icontains=q),
                    Q(client__code__iexact=q),
                    phone_q(q, "client__phone"),
                )
            )
            .select_related("client")
            .order_by("ready_at", "id")
        )
        if order is not None:
            cargos = cargos.exclude(order_lines__order=order)
        cargos = list(cargos[:CANDIDATE_LIMIT])
    return render(
        request,
        "panel/order/_candidates.html",
        {"cargos": cargos, "q": q, "day": day, "limit": CANDIDATE_LIMIT},
    )


# ---------- lines (HTMX: the row is swapped, counters out of band) ----------


@login_required
def line_row(request, day, pk):
    line = _get_line(_parse_day(day), pk)
    editing = request.GET.get("form") == CallStatus.REACHED
    context = _row_context(request, line, editing=editing)
    context["oob"] = False
    return render(request, "panel/order/_row_response.html", context)


def _promised(text: str) -> Date | None:
    text = (text or "").strip()
    if not text:
        return None
    try:
        return Date.fromisoformat(text)
    except ValueError:
        raise OrderError("Неверная дата приезда") from None


@require_POST
@login_required
def line_mark(request, day, pk):
    line = _get_line(_parse_day(day), pk)
    toast, error = "", False
    try:
        orders.mark_line(
            line,
            request.POST.get("status", ""),
            request.user,
            promised_date=_promised(request.POST.get("promised_date")),
            comment=request.POST.get("comment", ""),
        )
    except OrderError as exc:
        toast, error = str(exc), True
    if not _is_htmx(request):
        if error:
            messages.error(request, toast)
        return redirect(_day_url(line.order.date))
    line = _get_line(line.order.date, pk)
    context = _row_context(request, line, toast=toast, error=error)
    return render(request, "panel/order/_row_response.html", context)


@require_POST
@manager_required
def line_remove(request, day, pk):
    line = _get_line(_parse_day(day), pk)
    try:
        orders.remove_line(line)
    except OrderError as exc:
        toast, error, deleted = str(exc), True, False
    else:
        toast, error, deleted = f"Строка {line.article} убрана из ордера", False, True
    if not _is_htmx(request):
        (messages.error if error else messages.success)(request, toast)
        return redirect(_day_url(line.order.date))
    context = _row_context(request, line, toast=toast, error=error, deleted=deleted)
    return render(request, "panel/order/_row_response.html", context)


# ---------- print and Excel ----------


@login_required
def order_print(request, day):
    day = _parse_day(day)
    order = get_object_or_404(DailyOrder, date=day)
    return render(
        request,
        "panel/order/print.html",
        {"order": order, "day": day, "lines": numbered_lines(order)},
    )


def _status_text(line: DailyOrderLine) -> str:
    return "Выдан" if line.is_picked_up else line.get_call_status_display()


XLSX_COLUMNS = [
    ("№", 6),
    ("Артикул", 18),
    ("Клиент", 26),
    ("Код клиента", 13),
    ("Телефон", 16),
    ("Вес, кг", 10),
    ("Мест", 7),
    ("К оплате", 12),
    ("Дней", 7),
    ("Статус", 16),
    ("Приедет", 12),
    ("Комментарий", 36),
]


@login_required
def order_export(request, day):
    day = _parse_day(day)
    order = get_object_or_404(DailyOrder, date=day)
    book = Workbook()
    sheet = book.active
    sheet.title = f"Ордер {day:%d.%m.%Y}"
    sheet.append([title for title, _width in XLSX_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")
    for index, (_title, width) in enumerate(XLSX_COLUMNS):
        sheet.column_dimensions[chr(ord("A") + index)].width = width
    sheet.freeze_panes = "A2"
    phone_column = [title for title, _ in XLSX_COLUMNS].index("Телефон") + 1
    for line in numbered_lines(order):
        client: Client = line.cargo.client
        sheet.append(
            [
                line.number,
                line.article,
                line.client_name,
                client.code,
                line.phone,
                line.weight_kg,
                line.places,
                line.amount_due,
                line.days_waiting,
                _status_text(line),
                line.promised_date,
                line.comment,
            ]
        )
        row = sheet.max_row
        sheet.cell(row=row, column=phone_column).number_format = "@"
        sheet.cell(row=row, column=2).number_format = "@"
        sheet.cell(row=row, column=11).number_format = "DD.MM.YYYY"
    buffer = BytesIO()
    book.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="order-{day:%Y-%m-%d}.xlsx"'
    return response
