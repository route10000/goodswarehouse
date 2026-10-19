"""Cargo: list with filters, summary and bulk actions; cargo card and edit; add form."""

from datetime import datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db.models import Count, Q, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from logistics import services
from logistics.models import Cargo, Truck
from logistics.services import InvalidTransition
from notify.models import Notification
from panel.forms import CargoEditForm, CargoForm
from panel.ops_data import BLOCKS
from panel.permissions import manager_required

from .common import any_of, paginate, phone_q
from .lookup import lookup_context
from .trucks import create_cargo

AGE_CAPTIONS = {"A": "до 2 дней", "B": "3–4 дня", "C": "5–6 дней", "D": "7+ дней"}
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLSX_COLUMNS = (
    ("Артикул", 16),
    ("Описание", 28),
    ("Клиент", 24),
    ("Код клиента", 12),
    ("Телефон", 16),
    ("Трек-номер", 20),
    ("Продавец", 20),
    ("Упаковка", 12),
    ("Хрупкое", 9),
    ("Мест", 8),
    ("Вес, кг", 10),
    ("Объём, м³", 10),
    ("К оплате, ₽", 12),
    ("Оплачен", 9),
    ("Статус", 16),
    ("Машина", 14),
    ("Принят", 12),
    ("Готов к выдаче", 14),
    ("Хранение до", 12),
    ("Выдан", 12),
    ("Заметка", 30),
)
HISTORY_LIMIT = 20


def cargo_search_q(text: str) -> Q:
    return any_of(
        Q(article__icontains=text),
        Q(track_number__icontains=text),
        Q(seller__icontains=text),
        Q(client__code__iexact=text),
        phone_q(text, "client__phone"),
    )


def _day_start(day):
    return timezone.make_aware(datetime.combine(day, time.min))


def age_q(key: str, today=None) -> Q:
    """Ready cargo whose local days on the warehouse fall into block ``key`` (A–D)."""
    today = today or timezone.localdate()
    q = Q(status=Cargo.Status.READY)
    previous_max = None
    for block, _caption, max_days in BLOCKS:
        if block == key:
            break
        previous_max = max_days
    else:
        raise ValueError(key)
    if max_days is not None:
        newest = Q(ready_at__gte=_day_start(today - timedelta(days=max_days)))
        if previous_max is None:  # block A: a missing date counts as 0 days
            newest |= Q(ready_at__isnull=True)
        q &= newest
    if previous_max is not None:
        q &= Q(ready_at__lt=_day_start(today - timedelta(days=previous_max)))
    return q


def list_filters(request) -> dict:
    """Validated list filters from the query string (unknown values are dropped)."""
    params = request.GET
    status = params.get("status", "")
    truck_id = params.get("truck", "")
    age = params.get("age", "")
    pack = params.get("pack", "")
    return {
        "status": status if status in Cargo.Status.values else "",
        "truck_id": int(truck_id) if truck_id.isdigit() else None,
        "q": params.get("q", "").strip(),
        "age": age if age in AGE_CAPTIONS else "",
        "pack": pack if pack in Cargo.Packaging.values else "",
        "fragile": params.get("fragile", "") == "1",
    }


def filtered_cargo(filters: dict):
    cargos = Cargo.objects.select_related("client", "truck").order_by("-created_at", "-id")
    if filters["status"]:
        cargos = cargos.filter(status=filters["status"])
    if filters["truck_id"]:
        cargos = cargos.filter(truck_id=filters["truck_id"])
    if filters["q"]:
        cargos = cargos.filter(cargo_search_q(filters["q"]))
    if filters["age"]:
        cargos = cargos.filter(age_q(filters["age"]))
    if filters["pack"]:
        cargos = cargos.filter(packaging=filters["pack"])
    if filters["fragile"]:
        cargos = cargos.filter(fragile=True)
    return cargos


def summary_of(cargos) -> dict:
    totals = cargos.order_by().aggregate(
        count=Count("id"), places=Sum("places"), weight=Sum("weight_kg"), amount=Sum("amount_due")
    )
    return {
        "count": totals["count"],
        "places": totals["places"] or 0,
        "weight": totals["weight"] or Decimal("0"),
        "amount": totals["amount"] or Decimal("0"),
    }


@login_required
def cargo_list(request):
    filters = list_filters(request)
    cargos = filtered_cargo(filters)

    by_status = dict(Cargo.objects.values_list("status").annotate(n=Count("id")))
    status_filters = [{"value": "", "label": "Все", "count": sum(by_status.values())}]
    status_filters += [
        {"value": value, "label": label, "count": by_status.get(value, 0)}
        for value, label in Cargo.Status.choices
    ]
    has_filters = any(filters[key] for key in ("status", "truck_id", "q", "age", "pack", "fragile"))
    return render(
        request,
        "panel/cargo/list.html",
        {
            "nav_active": "cargo",
            "page": paginate(request, cargos),
            "summary": summary_of(cargos),
            **filters,
            "has_filters": has_filters,
            "filters": status_filters,
            "ages": [
                {"value": key, "label": f"{key} · {text}"} for key, text in AGE_CAPTIONS.items()
            ],
            "packagings": Cargo.Packaging.choices,
            "trucks": Truck.objects.order_by("-created_at", "-id"),
            "loading_trucks": Truck.objects.filter(status=Truck.Status.LOADING).order_by(
                "departure_date", "number"
            ),
            "next_url": request.get_full_path(),
        },
    )


def _safe_next(request) -> str:
    target = request.POST.get("next", "")
    if target.startswith("/panel/cargo/") and url_has_allowed_host_and_scheme(
        target, allowed_hosts={request.get_host()}, require_https=request.is_secure()
    ):
        return target
    return reverse("panel:cargo")


@require_POST
@manager_required
def cargo_bulk_load(request):
    """«Погрузить в машину…»: selected received cargo into one loading truck, all or nothing."""
    back = _safe_next(request)
    ids = [pk for pk in request.POST.getlist("cargo") if pk.isdigit()]
    truck_id = request.POST.get("truck", "")
    if not ids:
        messages.error(request, "Выберите грузы для погрузки.")
        return redirect(back)
    truck = Truck.objects.filter(pk=truck_id).first() if truck_id.isdigit() else None
    if truck is None:
        messages.error(request, "Выберите машину на погрузке.")
        return redirect(back)
    try:
        count = services.assign_many(ids, truck)
    except InvalidTransition as exc:
        messages.error(request, f"{exc}. Ничего не погружено: выберите только принятые грузы.")
    else:
        messages.success(request, f"Погружено в {truck.number}: {count}.")
    return redirect(back)


def _local_date(value):
    return timezone.localdate(value) if value else None


@manager_required
def cargo_export(request):
    """Excel of the selected cargo (POST ``cargo``) or of the whole current filter (GET)."""
    selected = [pk for pk in request.POST.getlist("cargo") if pk.isdigit()]
    if request.method == "POST" and selected:
        cargos = (
            Cargo.objects.filter(pk__in=selected)
            .select_related("client", "truck")
            .order_by("-created_at", "-id")
        )
    else:
        cargos = filtered_cargo(list_filters(request))

    book = Workbook()
    sheet = book.active
    sheet.title = "Грузы"
    sheet.append([title for title, _width in XLSX_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")
    for index, (_title, width) in enumerate(XLSX_COLUMNS):
        sheet.column_dimensions[chr(ord("A") + index)].width = width
    sheet.freeze_panes = "A2"
    titles = [title for title, _ in XLSX_COLUMNS]
    text_columns = [titles.index(name) + 1 for name in ("Артикул", "Телефон", "Трек-номер")]
    date_columns = [
        titles.index(name) + 1 for name in ("Принят", "Готов к выдаче", "Хранение до", "Выдан")
    ]
    for cargo in cargos.iterator():
        sheet.append(
            [
                cargo.article,
                cargo.description,
                cargo.client.name,
                cargo.client.code,
                cargo.client.phone,
                cargo.track_number,
                cargo.seller,
                cargo.get_packaging_display(),
                "да" if cargo.fragile else "нет",
                cargo.places,
                cargo.weight_kg,
                cargo.volume_m3,
                cargo.amount_due,
                "да" if cargo.paid else "нет",
                cargo.get_status_display(),
                cargo.truck.number if cargo.truck else "",
                _local_date(cargo.created_at),
                _local_date(cargo.ready_at),
                cargo.free_storage_until,
                _local_date(cargo.picked_up_at),
                cargo.note,
            ]
        )
        row = sheet.max_row
        for column in text_columns:
            sheet.cell(row=row, column=column).number_format = "@"
        for column in date_columns:
            sheet.cell(row=row, column=column).number_format = "DD.MM.YYYY"
    buffer = BytesIO()
    book.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=XLSX_MIME)
    stamp = timezone.localdate().strftime("%Y-%m-%d")
    response["Content-Disposition"] = f'attachment; filename="cargo-{stamp}.xlsx"'
    return response


# ---------- cargo card ----------

_AFTER_LOADING = {Cargo.Status.LOADED, Cargo.Status.IN_TRANSIT, Cargo.Status.READY}


def cargo_history(cargo: Cargo) -> list[dict]:
    """The cargo's way from its timestamps: received → loaded → departed → ready → picked up.

    Loading has no timestamp of its own, so that step shows the truck only.
    """
    truck = cargo.truck
    status = cargo.status
    picked = status == Cargo.Status.PICKED_UP
    by = cargo.picked_up_by
    steps = [
        {"key": "received", "label": "Принят", "at": cargo.created_at, "done": True, "text": ""},
        {
            "key": "loaded",
            "label": "Погружен в машину",
            "at": None,
            "done": truck is not None and (status in _AFTER_LOADING or picked),
            "text": truck.number if truck else "",
            "truck": truck,
        },
        {
            "key": "departed",
            "label": "Машина отправлена",
            "at": truck.departed_at if truck else None,
            "done": bool(truck and truck.departed_at),
            "text": truck.route if truck and truck.departed_at else "",
        },
        {
            "key": "ready",
            "label": "Прибыл, готов к выдаче",
            "at": cargo.ready_at or (truck.arrived_at if truck else None),
            "done": status == Cargo.Status.READY or picked,
            "text": "",
        },
        {
            "key": "picked_up",
            "label": "Выдан",
            "at": cargo.picked_up_at,
            "done": picked,
            "text": (f"выдал {by.get_full_name() or by.get_username()}" if by else "")
            + (" · оплачен" if picked and cargo.paid else ""),
        },
    ]
    current = next((index for index, step in enumerate(steps) if not step["done"]), None)
    for index, step in enumerate(steps):
        step["current"] = index == current
    return steps


@login_required
def cargo_detail(request, pk):
    cargo = get_object_or_404(
        Cargo.objects.select_related("client", "truck", "picked_up_by"), pk=pk
    )
    about = Q(cargo=cargo)
    if cargo.truck_id:
        about |= Q(truck_id=cargo.truck_id)
    notifications = (
        Notification.objects.filter(about, client_id=cargo.client_id)
        .select_related("truck")
        .order_by("-created_at", "-id")[:HISTORY_LIMIT]
    )
    today = timezone.localdate()
    return render(
        request,
        "panel/cargo/detail.html",
        {
            "nav_active": "cargo",
            "cargo": cargo,
            "history": cargo_history(cargo),
            "notifications": notifications,
            "attempts": cargo.contact_attempts.select_related("staff").order_by(
                "-created_at", "-id"
            )[:HISTORY_LIMIT],
            "storage_overdue": bool(
                cargo.free_storage_until
                and cargo.free_storage_until < today
                and cargo.status == Cargo.Status.READY
            ),
        },
    )


@manager_required
def cargo_edit(request, pk):
    cargo = get_object_or_404(Cargo.objects.select_related("client", "truck"), pk=pk)
    form = CargoEditForm(request.POST or None, request.FILES or None, cargo=cargo)
    if request.method == "POST" and form.is_valid():
        changes = form.changes()
        old_photo = cargo.photo.name if "photo" in changes else ""
        storage = cargo.photo.storage
        try:
            services.update_cargo(cargo, **changes)
        except ValidationError as exc:
            for field, errors in exc.message_dict.items():
                form.add_error(field if field in form.fields else None, errors)
        else:
            if old_photo and old_photo != cargo.photo.name:
                storage.delete(old_photo)
            messages.success(request, f"Груз {cargo.article} сохранён.")
            return redirect("panel:cargo_card", pk=cargo.pk)
    return render(
        request,
        "panel/cargo/form.html",
        {"nav_active": "cargo", "form": form, "cargo": Cargo.objects.get(pk=pk)},
    )


@manager_required
def cargo_create(request):
    initial = {}
    if request.GET.get("truck", "").isdigit():
        initial["truck"] = int(request.GET["truck"])
    form = CargoForm(request.POST or None, request.FILES or None, with_truck=True, initial=initial)
    if request.method == "POST" and form.is_valid():
        try:
            cargo = create_cargo(form, form.cleaned_data.get("truck"))
        except InvalidTransition as exc:
            form.add_error("truck", f"{exc}. Выберите машину на погрузке")
        else:
            messages.success(request, f"Груз {cargo.article} добавлен.")
            return redirect("panel:cargo")
    return render(
        request,
        "panel/cargo/new.html",
        {
            "nav_active": "cargo",
            "form": form,
            "lookup": lookup_context(request.POST.get("phone", ""), request.POST.get("name", "")),
        },
    )
