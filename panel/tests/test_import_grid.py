"""Editable import grid and manual entry (SPEC §14.5)."""

import io
import re

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from logistics import services
from logistics.imports import CELL_KEYS
from logistics.models import Cargo, Client

pytestmark = pytest.mark.django_db

HX = {"HTTP_HX_REQUEST": "true"}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ALL_HEADERS = [
    "артикул",
    "телефон",
    "имя",
    "описание",
    "мест",
    "вес",
    "объём",
    "сумма",
    "трек",
    "продавец",
    "упаковка",
    "хрупкое",
    "заметка",
]


def url(truck):
    return f"/panel/trucks/{truck.pk}/import/"


def grid(*rows, **action):
    """POST data of the grid: one list per column, in row order."""
    data = {key: [row.get(key, "") for row in rows] for key in CELL_KEYS}
    data.update(action)
    return data


def upload(rows, headers=ALL_HEADERS):
    wb = Workbook()
    wb.active.append(headers)
    for row in rows:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile("manifest.xlsx", buf.getvalue(), content_type=XLSX)


def article_inputs(page):
    return re.findall(r'name="article" value="([^"]*)"', page)


GOOD = {"article": "M-1", "phone": "8 914 123 45 67", "places": "2", "weight": "3,5"}
OTHER = {"article": "M-2", "phone": "89140000001", "name": "Ольга", "packaging": "bag"}


# ---------- manual entry ----------


def test_manual_entry_opens_empty_grid_with_five_rows(mgr, truck):
    response = mgr.get(url(truck), {"manual": "1"})

    page = response.content.decode()
    assert response.status_code == 200
    assert article_inputs(page) == [""] * 5
    assert "+ строка" in page
    assert "Сохранить 0 грузов" in page
    assert "Проверить" in page
    assert "Нет артикула" not in page  # an empty grid is not «checked» yet


def test_manual_entry_saves_rows_and_skips_empty_ones(mgr, truck, ivan):
    data = grid(
        {**GOOD, "track_number": "YT1", "seller": "1688", "fragile": "да", "note": "Стекло"},
        {"packaging": "box", "fragile": ""},  # untouched row
        OTHER,
        save="1",
        mode="manual",
    )

    response = mgr.post(url(truck), data, follow=True)

    assert response.redirect_chain[-1][0] == f"/panel/trucks/{truck.pk}/"
    first = Cargo.objects.get(article="M-1")
    assert first.client == ivan
    assert first.truck == truck and first.status == Cargo.Status.LOADED
    assert (first.places, str(first.weight_kg)) == (2, "3.50")
    assert (first.track_number, first.seller, first.fragile, first.note) == (
        "YT1",
        "1688",
        True,
        "Стекло",
    )
    second = Cargo.objects.get(article="M-2")
    assert second.packaging == "bag"
    assert Client.objects.get(phone="+79140000001").name == "Ольга"
    assert any("Сохранено 2 груза" in str(m) for m in response.context["messages"])


def test_check_shows_errors_under_the_row_and_saves_nothing(mgr, truck):
    data = grid(GOOD, {"article": "M-2", "phone": "12345", "places": "ноль"}, check="1")

    response = mgr.post(url(truck), data)

    page = response.content.decode()
    assert Cargo.objects.count() == 0
    assert "Неверный телефон: 12345" in page
    assert "Мест должно быть целым числом от 1" in page
    assert "grid-row__errors" in page
    assert 'name="phone" value="12345"' in page  # what was typed stays in the cell
    assert re.search(r'is-invalid[^>]*name="phone" value="12345"', page)
    assert article_inputs(page) == ["M-1", "M-2"]
    assert "2 строки, 1 с ошибками" in page


def test_fix_the_error_and_save(mgr, truck):
    bad = {"article": "M-2", "phone": "12345"}
    mgr.post(url(truck), grid(GOOD, bad, check="1"))
    assert Cargo.objects.count() == 0

    response = mgr.post(
        url(truck), grid(GOOD, {**bad, "phone": "+7 914 555-44-33"}, save="1"), follow=True
    )

    assert response.redirect_chain[-1][0] == f"/panel/trucks/{truck.pk}/"
    assert set(truck.cargos.values_list("article", flat=True)) == {"M-1", "M-2"}


def test_save_with_errors_saves_nothing(mgr, truck):
    response = mgr.post(
        url(truck), grid(GOOD, {"article": "M-1", "phone": "89140000002"}, save="1")
    )

    assert response.status_code == 200
    assert Cargo.objects.count() == 0
    assert Client.objects.count() == 0
    assert "Артикул повторяется (строка 1)" in response.content.decode()


def test_save_of_an_empty_grid_asks_for_a_row(mgr, truck):
    response = mgr.post(url(truck), grid({}, {}, save="1"))

    assert "Заполните хотя бы одну строку" in response.content.decode()
    assert Cargo.objects.count() == 0


def test_add_row_keeps_typed_values(mgr, truck):
    response = mgr.post(url(truck), grid(GOOD, OTHER, add_row="1"))

    page = response.content.decode()
    assert article_inputs(page) == ["M-1", "M-2", ""]
    assert Cargo.objects.count() == 0


def test_delete_row(mgr, truck):
    response = mgr.post(url(truck), grid(GOOD, {"article": "DROP"}, OTHER, delete="1"))

    page = response.content.decode()
    assert article_inputs(page) == ["M-1", "M-2"]


def test_delete_the_last_row_leaves_one_empty_row(mgr, truck):
    page = mgr.post(url(truck), grid(GOOD, delete="0")).content.decode()

    assert article_inputs(page) == [""]


def test_clean_check_keeps_preview_in_session(mgr, truck):
    mgr.post(url(truck), grid(GOOD, check="1"))

    stored = mgr.session[f"manifest:{truck.pk}"]
    assert [row["article"] for row in stored["rows"]] == ["M-1"]


def test_htmx_check_returns_only_the_grid(mgr, truck):
    response = mgr.post(url(truck), grid(GOOD, check="1"), **HX)

    page = response.content.decode()
    assert page.lstrip().startswith("<form")
    assert 'id="import-grid"' in page
    assert "<html" not in page


def test_htmx_save_redirects_the_page(mgr, truck):
    response = mgr.post(url(truck), grid(GOOD, save="1"), **HX)

    assert response["HX-Redirect"] == f"/panel/trucks/{truck.pk}/"
    assert Cargo.objects.filter(article="M-1").exists()


def test_grid_refused_for_truck_that_left(mgr, truck):
    services.truck_depart(truck)

    response = mgr.post(url(truck), grid(GOOD, save="1"), follow=True)

    assert Cargo.objects.count() == 0
    assert response.redirect_chain[-1][0] == f"/panel/trucks/{truck.pk}/"


# ---------- Excel file into the grid ----------


def test_uploaded_file_becomes_an_editable_grid_with_optional_columns(mgr, truck):
    rows = [
        [
            "X-1",
            "89141234567",
            "Иван",
            "Куртки",
            2,
            12.5,
            None,
            1500,
            "YT9",
            "Taobao",
            "Мешок",
            "да",
            "Осторожно",
        ],
        ["X-2", "123", "", "", 1, None, None, None, None, None, None, None, None],
    ]

    response = mgr.post(url(truck), {"file": upload(rows)})

    page = response.content.decode()
    assert article_inputs(page) == ["X-1", "X-2"]
    assert 'name="track_number" value="YT9"' in page
    assert 'name="seller" value="Taobao"' in page
    assert re.search(r'<option value="bag" selected>', page)
    assert re.search(r'<option value="да" selected>', page)
    assert "Неверный телефон: 123" in page


def test_upload_without_optional_columns_still_works(mgr, truck):
    headers = ALL_HEADERS[:8]
    response = mgr.post(
        url(truck), {"file": upload([["X-1", "89141234567", "", "", 1, None, None, None]], headers)}
    )

    page = response.content.decode()
    assert article_inputs(page) == ["X-1"]
    assert "Сохранить 1 груз" in page


def test_header_names_are_case_insensitive(mgr, truck):
    headers = [h.upper() for h in ALL_HEADERS]
    rows = [["X-1", "89141234567", "", "", 1, None, None, None, "YT9", "", "паллета", "", ""]]

    mgr.post(url(truck), {"file": upload(rows, headers)})
    mgr.post(
        url(truck),
        grid(
            {
                "article": "X-1",
                "phone": "89141234567",
                "track_number": "YT9",
                "packaging": "pallet",
            },
            save="1",
        ),
    )

    cargo = Cargo.objects.get(article="X-1")
    assert (cargo.track_number, cargo.packaging) == ("YT9", "pallet")


def test_template_has_all_columns(mgr):
    response = mgr.get("/panel/import/template.xlsx")

    wb = load_workbook(io.BytesIO(response.content))
    assert [c.value for c in wb.active[1]] == ALL_HEADERS


def test_import_page_mentions_optional_columns(mgr, truck):
    page = mgr.get(url(truck)).content.decode()

    assert "трек, продавец, упаковка, хрупкое, заметка" in page
    assert "Ввести вручную" in page
