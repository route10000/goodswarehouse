import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from logistics import services
from logistics.models import Cargo, Client

pytestmark = pytest.mark.django_db

HEADERS = ["артикул", "телефон", "имя", "описание", "мест", "вес", "объём", "сумма"]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def upload(rows, name="manifest.xlsx"):
    wb = Workbook()
    wb.active.append(HEADERS)
    for row in rows:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(name, buf.getvalue(), content_type=XLSX)


GOOD = [
    ["A-1", "89141234567", "Иван", "Куртки", 2, "12,5", None, 1500],
    ["A-2", "89140000001", "Новый", "", 1, None, None, None],
]


def test_upload_page(mgr, truck):
    response = mgr.get(f"/panel/trucks/{truck.pk}/import/")

    assert response.status_code == 200
    assert "Скачать шаблон" in response.content.decode()


def test_upload_shows_preview_and_stores_rows_in_session(mgr, truck, ivan):
    response = mgr.post(f"/panel/trucks/{truck.pk}/import/", {"file": upload(GOOD)})

    assert response.status_code == 200
    page = response.content.decode()
    assert "2 строки, 0 с ошибками, 1 новый клиент" in page
    assert "Сохранить 2 груза" in page
    stored = mgr.session[f"manifest:{truck.pk}"]
    assert [row["article"] for row in stored["rows"]] == ["A-1", "A-2"]
    assert Cargo.objects.count() == 0


def test_preview_with_errors_highlights_rows_and_hides_import(mgr, truck):
    rows = [*GOOD, ["", "123", "", "", 0, None, None, None]]

    response = mgr.post(f"/panel/trucks/{truck.pk}/import/", {"file": upload(rows)})

    page = response.content.decode()
    assert "3 строки, 1 с ошибками, 2 новых клиента" in page
    assert "row--error" in page
    assert "Нет артикула" in page
    assert 'value="apply"' not in page
    assert f"manifest:{truck.pk}" not in mgr.session


def test_apply_imports_from_session(mgr, truck, ivan):
    mgr.post(f"/panel/trucks/{truck.pk}/import/", {"file": upload(GOOD)})

    response = mgr.post(f"/panel/trucks/{truck.pk}/import/apply/", follow=True)

    assert response.redirect_chain[-1][0] == f"/panel/trucks/{truck.pk}/"
    assert Cargo.objects.filter(truck=truck).count() == 2
    assert Client.objects.filter(phone="+79140000001").exists()
    assert any("Импортировано 2" in str(m) for m in response.context["messages"])
    assert f"manifest:{truck.pk}" not in mgr.session


def test_apply_without_preview_does_nothing(mgr, truck):
    response = mgr.post(f"/panel/trucks/{truck.pk}/import/apply/", follow=True)

    assert Cargo.objects.count() == 0
    assert any(m.level_tag == "error" for m in response.context["messages"])


def test_apply_after_truck_left_rolls_back(mgr, truck):
    mgr.post(f"/panel/trucks/{truck.pk}/import/", {"file": upload(GOOD)})
    services.truck_depart(truck)

    response = mgr.post(f"/panel/trucks/{truck.pk}/import/apply/", follow=True)

    assert Cargo.objects.count() == 0
    assert not Client.objects.filter(phone="+79140000001").exists()
    assert any(m.level_tag == "error" for m in response.context["messages"])


def test_rejects_non_xlsx(mgr, truck):
    bad = SimpleUploadedFile("manifest.csv", b"a,b", content_type="text/csv")

    response = mgr.post(f"/panel/trucks/{truck.pk}/import/", {"file": bad})

    assert "Нужен файл .xlsx" in response.content.decode()
    assert f"manifest:{truck.pk}" not in mgr.session


def test_rejects_large_file(mgr, truck):
    big = SimpleUploadedFile("big.xlsx", b"0" * (2 * 1024 * 1024 + 1), content_type=XLSX)

    response = mgr.post(f"/panel/trucks/{truck.pk}/import/", {"file": big})

    assert "не больше 2 МБ" in response.content.decode()


def test_unreadable_xlsx_shows_file_error(mgr, truck):
    junk = SimpleUploadedFile("junk.xlsx", b"not a zip", content_type=XLSX)

    response = mgr.post(f"/panel/trucks/{truck.pk}/import/", {"file": junk})

    assert "Не удалось прочитать файл" in response.content.decode()


def test_import_only_for_loading_truck(mgr, truck):
    services.truck_depart(truck)

    response = mgr.get(f"/panel/trucks/{truck.pk}/import/", follow=True)

    assert response.redirect_chain[-1][0] == f"/panel/trucks/{truck.pk}/"
    assert any(m.level_tag == "error" for m in response.context["messages"])


def test_template_download(mgr):
    response = mgr.get("/panel/import/template.xlsx")

    assert response.status_code == 200
    assert response["Content-Type"] == XLSX
    assert "attachment" in response["Content-Disposition"]
    wb = load_workbook(io.BytesIO(response.content))
    assert [c.value for c in wb.active[1]][: len(HEADERS)] == HEADERS


def test_preview_shows_invalid_phone_as_typed(mgr, truck):
    rows = [["A", "12345", "", "", 1, None, None, None]]

    page = mgr.post(f"/panel/trucks/{truck.pk}/import/", {"file": upload(rows)}).content.decode()

    assert "+7 12345" not in page
    assert 'name="phone" value="12345"' in page
