from datetime import timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from logistics import services
from logistics.models import Cargo, ContactAttempt, DailyOrder, DailyOrderLine
from logistics.orders import add_line, build_daily_order, mark_line
from panel.tests.helpers import page_without_csrf

# «Today» is taken here and again by the views: pin the clock so midnight cannot split them.
pytestmark = [pytest.mark.django_db, pytest.mark.usefixtures("pinned_now")]

L = DailyOrderLine.CallStatus
HX = {"HTTP_HX_REQUEST": "true"}


def today():
    return timezone.localdate()


@pytest.fixture
def arrived(loaded_truck):
    services.truck_depart(loaded_truck)
    services.truck_arrive(loaded_truck)
    return loaded_truck


@pytest.fixture
def order(arrived, manager):
    """Today's order with all three ready cargos: B1 (Пётр, no Telegram) + A1, A2 added by hand."""
    order, _ = build_daily_order(today(), manager)
    for article in ("A1", "A2"):
        add_line(order, Cargo.objects.get(article=article), manager)
    return order


def line(article):
    return DailyOrderLine.objects.get(article=article)


def url(day=None, suffix=""):
    return f"/panel/order/{day or today():%Y-%m-%d}/{suffix}"


# ---------- page, empty state, build ----------


def test_today_without_order_shows_empty_state(mgr, arrived):
    page = mgr.get("/panel/order/", follow=True).content.decode()
    assert "Ордер на сегодня ещё не сформирован" in page
    assert "Сформировать ордер" in page


def test_build_creates_order_and_toasts_added_lines(mgr, arrived):
    response = mgr.post(url(suffix="build/"), follow=True)
    page = response.content.decode()
    assert DailyOrder.objects.filter(date=today()).exists()
    assert "Добавлено 1 строка" in page
    assert "B1" in page
    # A second build adds nothing.
    page = mgr.post(url(suffix="build/"), follow=True).content.decode()
    assert "Добавлено 0 строк" in page


def test_warehouse_can_build_and_open(wh, arrived):
    assert wh.post(url(suffix="build/")).status_code == 302
    assert wh.get("/panel/order/", follow=True).status_code == 200


def test_header_and_columns(mgr, order, petr):
    page = mgr.get("/panel/order/", follow=True).content.decode()
    assert f"Ордер на {today().day}" in page
    assert 'href="tel:+79141234568"' in page
    assert petr.code in page
    for column in ("Артикул", "Клиент", "Телефон", "Вес, кг", "Мест", "К оплате", "Статус"):
        assert f">{column}<" in page
    assert "Обновить ордер" in page and "Печать" in page and "Excel" in page


def test_bad_date_is_404(mgr):
    assert mgr.get("/panel/order/2026-13-45/").status_code == 404


def test_date_picker_redirects_to_day(mgr):
    response = mgr.get("/panel/order/?date=2026-10-01")
    assert response.status_code == 302
    assert response["Location"] == "/panel/warehouse/?tab=order&date=2026-10-01"


# ---------- counters and filters ----------


def counters(page):
    start = page.index('id="order-counters"')
    return page[start : page.index("</nav>", start)]


def test_counters_and_filters(mgr, order, manager):
    mark_line(line("A1"), L.REACHED, manager)
    mark_line(line("A2"), L.NO_ANSWER, manager)
    page = mgr.get("/panel/order/", follow=True).content.decode()
    block = counters(page)
    for label, n in (("Всего", 3), ("Позвонили", 1), ("Не дозвонились", 1), ("Осталось", 1)):
        assert f'{label} <span class="mono">{n}</span>' in block, label

    page = mgr.get("/panel/order/?status=pending", follow=True).content.decode()
    assert f'id="line-{line("B1").pk}"' in page
    assert f'id="line-{line("A1").pk}"' not in page

    page = mgr.get("/panel/order/?status=reached", follow=True).content.decode()
    assert f'id="line-{line("A1").pk}"' in page
    assert f'id="line-{line("B1").pk}"' not in page


def test_picked_up_row_goes_to_bottom_without_buttons(mgr, order, manager):
    a1 = line("A1")
    services.cargo_pickup(a1.cargo, manager, paid=True)
    page = mgr.get("/panel/order/", follow=True).content.decode()
    row_start = page.index(f'id="line-{a1.pk}"')
    row = page[row_start : page.index("</tr>", row_start)]
    assert "Выдан" in row
    assert "Позвонили" not in row
    assert page.index(f'id="line-{line("A2").pk}"') < row_start
    # A picked-up line is not «осталось».
    assert 'Осталось <span class="mono">2</span>' in counters(page)


# ---------- marking ----------


def test_mark_reached_with_promised_date(wh, order, keeper):
    a1 = line("A1")
    promised = today() + timedelta(days=1)
    response = wh.post(
        url(suffix=f"line/{a1.pk}/mark/"),
        {"status": "reached", "promised_date": promised.isoformat(), "comment": "вечером"},
        **HX,
    )
    page = response.content.decode()
    a1.refresh_from_db()
    assert a1.call_status == L.REACHED
    assert a1.promised_date == promised
    assert a1.comment == "вечером"
    attempt = ContactAttempt.objects.get()
    assert attempt.staff == keeper and attempt.promised_date == promised
    assert f'id="line-{a1.pk}"' in page
    assert f"приедет {promised:%d.%m}" in page


def test_mark_reached_without_date(mgr, order):
    a1 = line("A1")
    mgr.post(url(suffix=f"line/{a1.pk}/mark/"), {"status": "reached", "promised_date": ""}, **HX)
    a1.refresh_from_db()
    assert a1.call_status == L.REACHED and a1.promised_date is None


def test_mark_no_answer_with_comment(mgr, order):
    b1 = line("B1")
    mgr.post(url(suffix=f"line/{b1.pk}/mark/"), {"status": "no_answer", "comment": "сбросил"}, **HX)
    b1.refresh_from_db()
    assert b1.call_status == L.NO_ANSWER and b1.comment == "сбросил"


def test_correcting_a_mark(mgr, order):
    b1 = line("B1")
    mgr.post(url(suffix=f"line/{b1.pk}/mark/"), {"status": "no_answer"}, **HX)
    page = mgr.post(url(suffix=f"line/{b1.pk}/mark/"), {"status": "reached"}, **HX)
    b1.refresh_from_db()
    assert b1.call_status == L.REACHED
    assert ContactAttempt.objects.count() == 2
    # The buttons stay to correct the mark again.
    assert "Не дозвонились</button>" in page.content.decode()


def test_mark_response_updates_counters_and_badge_out_of_band(mgr, order):
    b1 = line("B1")
    page = mgr.post(url(suffix=f"line/{b1.pk}/mark/"), {"status": "reached"}, **HX)
    page = page.content.decode()
    assert 'id="order-counters"' in page and 'hx-swap-oob="true"' in page
    assert 'Позвонили <span class="mono">1</span>' in counters(page)
    assert 'id="order-badge"' in page


def test_reached_form_is_revealed_inline(mgr, order):
    b1 = line("B1")
    page = mgr.get(url(suffix=f"line/{b1.pk}/?form=reached"), **HX).content.decode()
    assert 'name="promised_date"' in page and "Сохранить" in page


def test_mark_error_returns_error_toast(mgr, order):
    b1 = line("B1")
    page = mgr.post(
        url(suffix=f"line/{b1.pk}/mark/"),
        {"status": "no_answer", "promised_date": today().isoformat()},
        **HX,
    ).content.decode()
    assert "Дату приезда можно указать только" in page
    assert "toast st-error" in page
    b1.refresh_from_db()
    assert b1.call_status == L.PENDING


def test_line_url_must_match_order_date(mgr, order):
    b1 = line("B1")
    other = today() + timedelta(days=1)
    assert mgr.post(url(other, f"line/{b1.pk}/mark/"), {"status": "reached"}).status_code == 404


# ---------- past orders ----------


@pytest.fixture
def past(order):
    yesterday = today() - timedelta(days=1)
    DailyOrder.objects.filter(pk=order.pk).update(date=yesterday)
    return yesterday


def test_past_order_is_read_only(mgr, past):
    page = mgr.get(url(past), follow=True).content.decode()
    assert "Ордер за прошедший день — только просмотр" in page
    assert "Позвонили</button>" not in page
    assert "Обновить ордер" not in page
    assert "Убрать" not in page


def test_past_order_mutations_fail(mgr, past):
    b1 = line("B1")
    page = mgr.post(url(past, f"line/{b1.pk}/mark/"), {"status": "reached"}, **HX)
    assert "только для чтения" in page.content.decode()
    page = mgr.post(url(past, f"line/{b1.pk}/remove/"), **HX)
    assert "только для чтения" in page.content.decode()
    page = mgr.post(url(past, "build/"), follow=True)
    assert "только для чтения" in page.content.decode()
    ready = Cargo.objects.filter(status=Cargo.Status.READY).first()
    page = mgr.post(url(past, "add/"), {"cargo": ready.pk}, follow=True)
    assert "только для чтения" in page.content.decode()
    b1.refresh_from_db()
    assert b1.call_status == L.PENDING
    assert DailyOrderLine.objects.count() == 3
    assert not ContactAttempt.objects.exists()


# ---------- manager-only remove / add ----------


def test_manager_removes_pending_line(mgr, order):
    b1 = line("B1")
    page = mgr.post(url(suffix=f"line/{b1.pk}/remove/"), **HX).content.decode()
    assert not DailyOrderLine.objects.filter(pk=b1.pk).exists()
    assert 'id="order-counters"' in page


def test_remove_marked_line_shows_error(mgr, order, manager):
    b1 = mark_line(line("B1"), L.REACHED, manager)
    page = mgr.post(url(suffix=f"line/{b1.pk}/remove/"), **HX).content.decode()
    assert "убрать нельзя" in page
    assert DailyOrderLine.objects.filter(pk=b1.pk).exists()


def test_warehouse_cannot_remove_or_add(wh, order):
    b1 = line("B1")
    assert wh.post(url(suffix=f"line/{b1.pk}/remove/"), **HX).status_code == 403
    assert wh.post(url(suffix="add/"), {"cargo": b1.cargo_id}).status_code == 403
    assert wh.get(url(suffix="candidates/?q=A1")).status_code == 403
    page = wh.get("/panel/order/", follow=True).content.decode()
    assert "Убрать" not in page and "Добавить груз" not in page


def test_manager_searches_and_adds_cargo(mgr, arrived, manager):
    order, _ = build_daily_order(today(), manager)
    page = page_without_csrf(mgr.get(url(suffix="candidates/?q=A1"), **HX))
    assert "A1" in page and "B1" not in page  # B1 is already in the order
    page = mgr.get(url(suffix="candidates/?q=4567"), **HX).content.decode()
    assert "A1" in page and "A2" in page  # by phone of Иван
    a1 = Cargo.objects.get(article="A1")
    page = mgr.post(url(suffix="add/"), {"cargo": a1.pk}, follow=True).content.decode()
    assert order.lines.filter(cargo=a1).exists()
    assert "добавлен в ордер" in page
    page = mgr.post(url(suffix="add/"), {"cargo": a1.pk}, follow=True).content.decode()
    assert "уже есть в ордере" in page


# ---------- print and Excel ----------


def test_print_view(mgr, order, manager):
    mark_line(line("A1"), L.REACHED, manager, promised_date=today() + timedelta(days=1))
    page = mgr.get(url(suffix="print/")).content.decode()
    assert f"Ордер на обзвон · {today():%d.%m.%Y} · GoodsWarehouse" in page
    assert "Дозвонился" in page and "Не дозвонился" in page and "Комментарий" in page
    assert "window.print()" in page
    assert "B1" in page and "A2" in page
    assert "☒" in page  # the reached mark
    assert 'class="sidebar"' not in page


def test_xlsx_export(mgr, order, manager, petr):
    mark_line(line("A2"), L.NO_ANSWER, manager, comment="занято")
    response = mgr.get(url(suffix="export.xlsx"))
    assert response.status_code == 200
    assert f'filename="order-{today():%Y-%m-%d}.xlsx"' in response["Content-Disposition"]
    sheet = load_workbook(BytesIO(response.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    header = rows[0]
    assert header[:4] == ("№", "Артикул", "Клиент", "Код клиента")
    assert sheet.cell(row=1, column=1).font.bold
    by_article = {row[1]: dict(zip(header, row, strict=True)) for row in rows[1:]}
    b1 = by_article["B1"]
    assert b1["Телефон"] == "+79141234568"
    assert b1["Код клиента"] == petr.code
    assert b1["Статус"] == "Не звонили"
    assert b1["Мест"] == 3
    assert Decimal(str(b1["К оплате"])) == Decimal("50")
    assert by_article["A2"]["Статус"] == "Не дозвонились"
    assert by_article["A2"]["Комментарий"] == "занято"
    phone_col = header.index("Телефон") + 1
    assert sheet.cell(row=2, column=phone_col).number_format == "@"
    assert sheet.column_dimensions["B"].width > 10


# ---------- sidebar badge ----------


def test_sidebar_badge_counts_remaining_lines(mgr, order, manager):
    mark_line(line("A1"), L.REACHED, manager)
    page = mgr.get("/panel/trucks/").content.decode()
    assert 'id="order-badge">2<' in page


def test_sidebar_badge_hidden_when_nothing_left(mgr, arrived):
    page = mgr.get("/panel/trucks/").content.decode()
    assert 'id="order-badge" hidden' in page
