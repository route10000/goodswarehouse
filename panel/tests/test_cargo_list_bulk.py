"""Cargo list: summary, age / packaging filters, track search, bulk load and Excel export."""

import io
from datetime import timedelta
from decimal import Decimal

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from logistics import services
from logistics.models import Cargo

pytestmark = pytest.mark.django_db


def articles(response):
    return sorted(c.article for c in response.context["page"].object_list)


@pytest.fixture
def stock(ivan, petr, pinned_now):
    """Ready cargo with 0, 3, 5 and 9 local days on the warehouse plus two received ones."""
    made = {}
    for article, days in (("AGE-A", 0), ("AGE-B", 3), ("AGE-C", 5), ("AGE-D", 9)):
        truck = services.create_truck()
        services.add_cargo(ivan, article, truck=truck, places=2, weight_kg=Decimal("1.5"))
        services.truck_depart(truck, now=pinned_now - timedelta(days=days, hours=1))
        services.truck_arrive(truck, now=pinned_now - timedelta(days=days))
    made["r1"] = services.add_cargo(
        petr,
        "REC-1",
        places=4,
        weight_kg=Decimal("10"),
        amount_due=Decimal("300"),
        track_number="YT-TRACK-77",
        seller="Lucky Store",
        packaging=Cargo.Packaging.PALLET,
        fragile=True,
    )
    made["r2"] = services.add_cargo(petr, "REC-2", places=1, amount_due=Decimal("200.50"))
    return made


# ---------- summary, filters, search ----------


def test_summary_for_current_filter(mgr, stock):
    summary = mgr.get("/panel/cargo/", {"status": "received"}).context["summary"]

    assert summary["count"] == 2
    assert summary["places"] == 5
    assert summary["weight"] == Decimal("10")
    assert summary["amount"] == Decimal("500.50")
    page = mgr.get("/panel/cargo/", {"status": "received"}).content.decode()
    assert "Мест" in page and "К оплате" in page


def test_summary_counts_whole_filter_not_page(mgr, ivan):
    for n in range(55):
        services.add_cargo(ivan, f"P{n:02d}", places=2)

    assert mgr.get("/panel/cargo/").context["summary"]["places"] == 110


@pytest.mark.parametrize(
    ("age", "expected"),
    [("A", ["AGE-A"]), ("B", ["AGE-B"]), ("C", ["AGE-C"]), ("D", ["AGE-D"])],
)
def test_age_filter(mgr, stock, age, expected):
    assert articles(mgr.get("/panel/cargo/", {"age": age})) == expected


def test_unknown_age_is_ignored(mgr, stock):
    assert len(articles(mgr.get("/panel/cargo/", {"age": "Z"}))) == 6


def test_packaging_and_fragile_filters(mgr, stock):
    assert articles(mgr.get("/panel/cargo/", {"pack": "pallet"})) == ["REC-1"]
    assert articles(mgr.get("/panel/cargo/", {"fragile": "1"})) == ["REC-1"]
    assert len(articles(mgr.get("/panel/cargo/", {"pack": "box"}))) == 5


def test_search_by_track_and_seller(mgr, stock):
    assert articles(mgr.get("/panel/cargo/", {"q": "track-77"})) == ["REC-1"]
    assert articles(mgr.get("/panel/cargo/", {"q": "lucky"})) == ["REC-1"]


def test_list_shows_track_and_row_links(mgr, stock):
    page = mgr.get("/panel/cargo/").content.decode()

    assert "YT-TRACK-77" in page
    assert f'data-href="/panel/cargo/{stock["r1"].pk}/"' in page


def test_empty_state_has_hint_and_action(mgr, db):
    page = mgr.get("/panel/cargo/").content.decode()

    assert "Грузов пока нет" in page
    assert "/panel/cargo/new/" in page


def test_bulk_controls_only_for_manager(mgr, wh, stock, truck):
    assert "data-bulk-item" in mgr.get("/panel/cargo/").content.decode()
    page = wh.get("/panel/cargo/").content.decode()
    assert "data-bulk-item" not in page
    assert "/panel/cargo/load/" not in page


# ---------- bulk load ----------


def test_bulk_load_success(mgr, stock, truck):
    ids = [stock["r1"].pk, stock["r2"].pk]
    response = mgr.post(
        "/panel/cargo/load/", {"cargo": ids, "truck": truck.pk, "next": "/panel/cargo/?q=REC"}
    )

    assert response.status_code == 302
    assert response["Location"] == "/panel/cargo/?q=REC"
    assert set(truck.cargos.values_list("article", flat=True)) == {"REC-1", "REC-2"}
    assert set(Cargo.objects.filter(pk__in=ids).values_list("status", flat=True)) == {"loaded"}


def test_bulk_load_mixed_statuses_rejected_atomically(mgr, stock, truck):
    ready = Cargo.objects.get(article="AGE-A")
    response = mgr.post(
        "/panel/cargo/load/", {"cargo": [stock["r1"].pk, ready.pk], "truck": truck.pk}, follow=True
    )

    assert "AGE-A нельзя погрузить" in response.content.decode()
    assert not truck.cargos.exists()
    stock["r1"].refresh_from_db()
    assert stock["r1"].status == Cargo.Status.RECEIVED


def test_bulk_load_non_loading_truck_rejected(mgr, stock, truck, ivan):
    services.add_cargo(ivan, "IN-TRUCK", truck=truck)
    services.truck_depart(truck)
    response = mgr.post(
        "/panel/cargo/load/", {"cargo": [stock["r1"].pk], "truck": truck.pk}, follow=True
    )

    assert "уже не грузится" in response.content.decode()
    stock["r1"].refresh_from_db()
    assert stock["r1"].status == Cargo.Status.RECEIVED


def test_bulk_load_requires_selection_and_truck(mgr, stock, truck):
    page = mgr.post("/panel/cargo/load/", {"truck": truck.pk}, follow=True).content.decode()
    assert "Выберите грузы" in page
    page = mgr.post("/panel/cargo/load/", {"cargo": [stock["r1"].pk]}, follow=True)
    assert "Выберите машину" in page.content.decode()


def test_bulk_load_ignores_unsafe_next(mgr, stock, truck):
    response = mgr.post(
        "/panel/cargo/load/",
        {"cargo": [stock["r1"].pk], "truck": truck.pk, "next": "https://evil.example/"},
    )
    assert response["Location"] == "/panel/cargo/"


def test_bulk_load_forbidden_for_warehouse(wh, stock, truck):
    response = wh.post("/panel/cargo/load/", {"cargo": [stock["r1"].pk], "truck": truck.pk})

    assert response.status_code == 403
    assert not truck.cargos.exists()


def test_bulk_load_get_not_allowed(mgr):
    assert mgr.get("/panel/cargo/load/").status_code == 405


# ---------- Excel export ----------


def read_xlsx(response):
    assert response.status_code == 200
    assert response["Content-Type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response["Content-Disposition"]
    sheet = load_workbook(io.BytesIO(response.content)).active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def test_export_whole_filter(mgr, stock):
    rows = read_xlsx(mgr.get("/panel/cargo/export.xlsx", {"status": "received"}))

    header = rows[0]
    for title in (
        "Артикул",
        "Клиент",
        "Код клиента",
        "Телефон",
        "Трек-номер",
        "Продавец",
        "Упаковка",
        "Хрупкое",
        "Мест",
        "Вес, кг",
        "Объём, м³",
        "К оплате, ₽",
        "Статус",
        "Машина",
        "Заметка",
    ):
        assert title in header, title
    body = {row[header.index("Артикул")]: row for row in rows[1:]}
    assert sorted(body) == ["REC-1", "REC-2"]
    rec = body["REC-1"]
    assert rec[header.index("Трек-номер")] == "YT-TRACK-77"
    assert rec[header.index("Продавец")] == "Lucky Store"
    assert rec[header.index("Упаковка")] == "Паллета"
    assert rec[header.index("Хрупкое")] == "да"
    assert rec[header.index("Телефон")] == "+79141234568"
    assert rec[header.index("Мест")] == 4
    assert rec[header.index("Статус")] == "Принят"


def test_export_selected_only(mgr, stock):
    response = mgr.post("/panel/cargo/export.xlsx", {"cargo": [stock["r2"].pk]})
    rows = read_xlsx(response)

    assert [row[0] for row in rows[1:]] == ["REC-2"]


def test_export_forbidden_for_warehouse(wh, stock):
    assert wh.get("/panel/cargo/export.xlsx").status_code == 403


def test_age_filter_uses_local_days(mgr, ivan, pinned_now):
    """A cargo that arrived yesterday late evening is 1 day old (block A), not 0 or 2."""
    truck = services.create_truck()
    services.add_cargo(ivan, "LATE", truck=truck)
    services.truck_depart(truck, now=pinned_now - timedelta(days=3))
    local = timezone.localtime(pinned_now)
    yesterday_evening = local.replace(hour=23, minute=0) - timedelta(days=1)
    services.truck_arrive(truck, now=yesterday_evening)

    assert articles(mgr.get("/panel/cargo/", {"age": "A"})) == ["LATE"]
